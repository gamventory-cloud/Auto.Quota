# table_guide.py
#
# 사내 '테이블 가이드' 엑셀과 .sav 로 Command.sps(배너·리코드) 와 Table.sps 를 만든다.
#
#   read_guide      가이드의 Basic Table / Banner / (있으면) TG 시트를 읽는다
#   infer_tables    가이드 한 줄마다 유형·변수·베이스를 추론한다 (상태 표시 포함)
#   infer_banners   Banner 시트로 bv/cv 배너를 만든다
#   write_command   3] Command.sps
#   write_table     4] Table.sps  — 배너를 행(ROW) 또는 열(COL) 에 둔다
#
# 추론은 규칙으로만 한다. 자신 없는 줄은 상태를 CHECK/NEED 로 남겨 사람이 고친다.
# 화면에서 고친 값(유형·변수·조건 등)을 그대로 받아 신텍스를 다시 만들 수 있게
# 표 정의(TableSpec)는 문자열 필드 위주로 둔다.

from __future__ import annotations

import collections
import io
import re
from dataclasses import dataclass, field

import pandas as pd

ROW, COL = "row", "col"               # 배너를 행에 / 열에

OK, CHECK, NEED = "ok", "check", "need"
STATUS_LABEL = {OK: "✅ 확인됨", CHECK: "⚠️ 확인 필요", NEED: "❌ 직접 입력"}

KINDS = ["single", "multi", "mean", "scale", "summary"]
KIND_LABEL = {"single": "단수", "multi": "복수", "mean": "수치형(평균)",
              "scale": "척도", "summary": "Summary"}
LABEL_KIND = {v: k for k, v in KIND_LABEL.items()}

_CODE = re.compile(r"\s*([A-Za-z]+\d+)((?:[-_]\d+)*)\s*[.)]?\s*(.*)", re.S)
_WORD = re.compile(r"[가-힣A-Za-z0-9%]+")


# =============================================================================
# 가이드 읽기
# =============================================================================
@dataclass
class GuideRow:
    row: int                  # 엑셀 줄 번호 (안내용)
    title: str
    extra_banner: str = ""
    base: str = ""
    stats: str = ""
    note: str = ""


@dataclass
class BannerRow:
    code: str                 # 가이드에 적힌 문항 (SQ1, B7-1 …)
    label: str                # 배너 이름
    values: list[tuple[int, str]] = field(default_factory=list)   # 적힌 보기


@dataclass
class Guide:
    rows: list[GuideRow]
    banners: list[BannerRow]
    tg: list[tuple] = field(default_factory=list)
    section_title: str = "응답자 특성"
    info: dict = field(default_factory=dict)


def _s(v) -> str:
    return "" if v is None else str(v).strip()


def _header_cols(ws, names: dict[str, tuple[str, ...]], max_scan: int = 15):
    """머리줄을 찾아 {키: 열 번호} 를 돌려준다. 조사마다 열 위치가 조금씩 다르다."""
    for r in ws.iter_rows(min_row=1, max_row=max_scan):
        found = {}
        for c in r:
            txt = _s(c.value).replace(" ", "").lower()
            for key, cands in names.items():
                if key not in found and any(txt.startswith(x) for x in cands):
                    found[key] = c.column - 1
        if "title" in found or "code" in found:
            return r[0].row, found
    return None, {}


def _parse_value_text(text: str) -> list[tuple[int, str]]:
    """'1) 남성  2) 여성' / 줄바꿈·탭으로 나뉜 보기 → [(1,'남성'), (2,'여성')]."""
    parts = re.split(r"(?:^|\s)(\d+)\)\s*", _s(text))
    out = []
    for i in range(1, len(parts) - 1, 2):
        lab = re.sub(r"\s+", " ", parts[i + 1]).strip()
        if lab:
            out.append((int(parts[i]), lab))
    return out


def read_guide(data: bytes) -> Guide:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), data_only=True)
    if "Basic Table" not in wb.sheetnames:
        raise ValueError("가이드에 'Basic Table' 시트가 없습니다.")
    ws = wb["Basic Table"]
    hdr, cols = _header_cols(ws, {
        "title": ("문항명(요약본)", "문항명(요약)"),
        "extra": ("추가배너",),
        "base": ("베이스",),
        "stats": ("통계값",),
        "note": ("비고",),
    })
    if "title" not in cols:
        raise ValueError("Basic Table 시트에서 '문항명 (요약본)' 머리줄을 찾지 못했습니다.")

    rows: list[GuideRow] = []
    section = ""
    for r in ws.iter_rows(min_row=hdr + 1):
        vals = [c.value for c in r]
        get = lambda k: _s(vals[cols[k]]) if k in cols and cols[k] < len(vals) else ""  # noqa: E731
        title = get("title")
        if not title:
            continue
        if not _CODE.match(title) or not re.match(r"\s*[A-Za-z]+\d", title):
            if not rows and not section and title.startswith("◈"):
                section = title.lstrip("◈ ").strip()          # '◈ 응답자 특성'
            continue
        rows.append(GuideRow(r[0].row, title, get("extra"), get("base"),
                             get("stats"), get("note")))

    banners: list[BannerRow] = []
    info = {}
    if "Banner" in wb.sheetnames:
        wsb = wb["Banner"]
        for r in wsb.iter_rows(min_row=1, max_row=12, values_only=True):
            cells = [_s(x) for x in r]
            for i, c in enumerate(cells[:-1]):
                if c.replace(" ", "") in ("KMSNo.", "KMSNo") and cells[i + 1]:
                    info["kms"] = cells[i + 1]
                if c == "프로젝트명" and cells[i + 1]:
                    info["project"] = cells[i + 1]
        bh, bcols = _header_cols(wsb, {"code": ("문항",), "label": ("variablelabel",),
                                       "values": ("valuelabel",)}, max_scan=20)
        if bh:
            for r in wsb.iter_rows(min_row=bh + 1, values_only=True):
                code = _s(r[bcols["code"]]) if "code" in bcols else ""
                if not re.match(r"[A-Za-z]+\d", code):
                    if banners:
                        break
                    continue
                label = _s(r[bcols.get("label", 1)])
                # 보기 칸 위치가 조사마다 C / D 로 달라서 머리줄 오른쪽에서 찾는다
                vtxt = ""
                start = bcols.get("values", 2)
                for j in range(start, min(start + 3, len(r))):
                    if _s(r[j]) and not _s(r[j]).isdigit():
                        vtxt = _s(r[j])
                        break
                banners.append(BannerRow(code, label, _parse_value_text(vtxt)))

    tg = []
    if "TG" in wb.sheetnames:
        tg = [tuple(r) for r in wb["TG"].iter_rows(min_row=2, values_only=True)]
    return Guide(rows, banners, tg, section or "응답자 특성", info)


# =============================================================================
# 데이터 쪽 도우미
# =============================================================================
class Data:
    """.sav 의 변수 정보를 대소문자 무관하게 찾기 쉽게 묶어 둔다."""

    def __init__(self, df: pd.DataFrame, meta):
        self.df = df
        self.cols = list(df.columns)
        self.up = {c.upper(): c for c in self.cols}
        self.pos = {c: i for i, c in enumerate(self.cols)}
        self.vl = meta.variable_value_labels
        self.cl = {c: _s(meta.column_names_to_labels.get(c)) for c in self.cols}

    def get(self, name: str):
        return self.up.get(str(name).replace("-", "_").upper())

    def family(self, code: str) -> list[str]:
        return [c for c in self.cols
                if re.fullmatch(rf"{re.escape(code)}(_\w+)?", c, re.I)
                and not c.lower().endswith("_etc")]

    def numbered(self, code: str) -> list[str]:
        """code_1, code_2 … (숫자 하나만 더 붙은 것)."""
        return [c for c in self.cols if re.fullmatch(rf"{re.escape(code)}_\d+", c, re.I)]

    def labels(self, v: str) -> dict:
        return self.vl.get(v, {}) or {}

    def has_data(self, v: str) -> bool:
        """빈칸·결측이 아닌 응답이 하나라도 있는지."""
        s = self.df[v]
        if s.dtype == object:
            return bool(s.astype(str).str.strip().replace({"nan": ""}).ne("").any())
        return bool(s.notna().any())

    def ma_set(self, code: str) -> list[str]:
        """code_1, code_2 … 가 복수응답 세트(변수마다 자기 코드 하나)면 그 목록."""
        from banner_table_engine import is_category_coded_set

        mem = [c for c in self.numbered(code) if self.labels(c)]
        if len(mem) < 2 or any(self.labels(c) != self.labels(mem[0]) for c in mem):
            return []
        return mem if is_category_coded_set(self.df, mem) else []

    def valid_codes(self, v: str) -> list[float]:
        """척도 계산에 쓸 보기 코드. 모름·무응답처럼 번호가 동떨어진 코드는 뺀다.

        '전혀 모름 / 잘 모름' 처럼 척도의 한 단계가 '모름' 으로 끝나기도 하므로
        라벨 문구로는 판단하지 않는다. 99 같은 큰 코드, 또는 1~5 뒤의 9 처럼
        이어지지 않는 마지막 코드만 뺀다.
        """
        lab = self.labels(v)
        if lab:
            keep = sorted(lab)
            if len(keep) > 2 and keep[-1] >= 90 and keep[-2] < 90:
                keep = keep[:-1]
            elif len(keep) > 3 and keep[-1] - keep[-2] > 1 and \
                    all(b - a == 1 for a, b in zip(keep[:-2], keep[1:-1])):
                keep = keep[:-1]
            return keep
        vals = pd.to_numeric(self.df[v], errors="coerce").dropna().unique().tolist()
        return sorted(vals)

    def var_list(self, vs: list[str]) -> str:
        """파일 순서로 이어지면 'a to b', 아니면 공백으로."""
        if len(vs) > 1:
            idx = [self.pos[v] for v in vs]
            if idx == list(range(idx[0], idx[0] + len(vs))):
                return f"{vs[0].lower()} to {vs[-1].lower()}"
        return " ".join(v.lower() for v in vs)


def parse_vars(text: str, data: Data) -> tuple[list[str], list[str]]:
    """'a to b' · 'a, b' · 'a b' → 변수 목록 (파일 순서). (목록, 문제)"""
    text = _s(text)
    if not text:
        return [], []
    out, bad = [], []
    for chunk in re.split(r"\s*,\s*", text):
        m = re.fullmatch(r"(\S+)\s+to\s+(\S+)", chunk, re.I)
        if m:
            a, b = data.get(m.group(1)), data.get(m.group(2))
            if a is None or b is None:
                bad.append(chunk)
                continue
            i0, i1 = sorted((data.pos[a], data.pos[b]))
            out += data.cols[i0:i1 + 1]
        else:
            for tok in chunk.split():
                v = data.get(tok)
                (out.append(v) if v else bad.append(tok))
    seen = set()
    return [v for v in out if not (v in seen or seen.add(v))], bad


# =============================================================================
# 표 정의 추론
# =============================================================================
@dataclass
class TableSpec:
    title: str
    kind: str                         # KINDS
    vars: str                         # 'q1' / 'q16_1 to q16_3' / 'b2_1 b2_2 …'
    cond: str                         # SPSS 조건 ('' = 전체)
    sort: bool = False
    extra_banner: str = ""            # 추가 배너 변수 (없으면 '')
    stats: str = ""                   # 가이드 통계값 (척도 묶음·100점 환산)
    base: str = ""                    # 가이드 베이스 원문 (안내용)
    status: str = OK
    reason: str = ""
    row: int = 0


def _overlap(a: str, b: str) -> float:
    wa, wb = set(_WORD.findall(a)), set(_WORD.findall(b))
    return len(wa & wb) / (len(wa) or 1)


def _item_text(label: str) -> str:
    """'B2. [진료항목의 실제 지불 비용] 1) (10kg 미만)초진진찰' → '1) (10kg 미만)초진진찰'."""
    s = re.sub(r"^\s*[A-Za-z]+\d+(?:[-_]\d+)*\.\s*", "", label)
    if "]" in s:
        s = s.rsplit("]", 1)[-1]
    return s.strip()


def _label_code(label: str) -> str:
    """라벨 맨 앞 문항 번호. 'Q12-1. 귀하께서는 …' → 'Q12_1', 없으면 ''."""
    m = re.match(r"\s*([A-Za-z]+\d+(?:-\d+)*)\.", label or "")
    return m.group(1).replace("-", "_").upper() if m else ""


def _norm(s: str) -> str:
    return re.sub(r"[^가-힣A-Za-z0-9%]", "", s).lower()


def _find_vars(title: str, data: Data) -> tuple[list[str], str | None, str, str]:
    """제목 → (변수들, 유형 힌트, 상태, 사유)."""
    m = _CODE.match(title)
    q, subs, rest = m.group(1), m.group(2), m.group(3)
    sub_dash = re.findall(r"-(\d+)", subs)       # 'Q27-1' 의 1  (하위 문항)
    sub_us = re.findall(r"_(\d+)", subs)         # 'B4_1' 의 1   (문항 안의 부분)
    code = q + subs.replace("-", "_")

    # 순위: '1순위' / '2순위' / '1~3순위' / '1+2순위'
    rk = re.search(r"(\d(?:\s*[+~]\s*\d)*)\s*순위", rest)
    if rk:
        nums = [int(x) for x in re.findall(r"\d", rk.group(1))]
        ranks = list(range(nums[0], nums[-1] + 1)) if "~" in rk.group(1) else nums
        kind = "single" if len(ranks) == 1 else "multi"
        mem = [c for c in data.cols if re.fullmatch(rf"{re.escape(code)}_\d", c, re.I)]
        if len(mem) >= max(ranks):
            return [mem[r - 1] for r in ranks], kind, OK, ""
        # 라벨의 'N순위' 로 찾는다. '[에너지]' 같은 꼬리표가 있으면 라벨에도 있어야 한다.
        tags = [_norm(t) for t in re.findall(r"\[([^\]]+)\]", rest)]
        by_rank: dict[int, list[str]] = collections.defaultdict(list)
        for c in data.family(code) or data.family(q):
            m2 = re.search(r"(\d)\s*순위", data.cl[c])
            # 꼬리표는 대괄호 뒤(항목 부분)에서만 본다. 대괄호 안 문항 문구에는
            # '(에너지 및 자동차 분야)' 처럼 모든 꼬리표가 다 들어 있기도 하다.
            item = _norm(_item_text(data.cl[c]))
            if m2 and data.labels(c) and all(t in item for t in tags):
                by_rank[int(m2.group(1))].append(c)
        if all(len(by_rank.get(r, [])) == 1 for r in ranks):
            return ([by_rank[r][0] for r in ranks], kind, CHECK,
                    "라벨의 'N순위'" + (" 와 꼬리표" if tags else "") + " 로 순위 변수를 찾았습니다")
        if sub_dash == ["1"]:                     # 'Q32-1' 의 순위가 Q32_1 ~ Q32_3 인 경우
            mem = [c for c in data.cols if re.fullmatch(rf"{re.escape(q)}_\d", c, re.I)
                   and data.labels(c) and data.has_data(c)]
            if len(mem) >= max(ranks):
                return ([mem[r - 1] for r in ranks], kind, CHECK,
                        f"{code}_1 이 없어 {mem[0]} ~ 을 순위 변수로 봤습니다")
        return [], None, NEED, f"순위 변수 {code}_1 … 을 찾지 못했습니다"

    # 항목 번호: '_3)' / ' - 3)' 또는 끝의 '_3' / '- 3'
    it = (re.search(r"_(\d+)\)", rest) or re.search(r"(?:^|\s)-\s*(\d+)\)", rest)
          or re.search(r"(?:_|\s-\s*)(\d+)\s*$", rest))
    k = it.group(1) if it else None
    names = []
    if k:
        if sub_us:                                # 'B4_1. …_3)' → B4_3_1
            mid = "_".join(sub_dash)
            names.append(f"{q}_{mid + '_' if mid else ''}{k}_{sub_us[-1]}")
        names.append(f"{code}_{k}")
        if sub_dash and not sub_us:               # 'Q27-1. …_2' → Q27_2_1
            names.append(f"{q}_{k}_{sub_dash[-1]}")
    for nm in names:
        v = data.get(nm)
        # 응답이 하나도 없는 변수는 건너뛴다. 'Q4-1. … - 1)' 의 Q4_1_1 은 항목이 아니라
        # 뒤따르는 주관식 문항이라 비어 있다.
        if v and data.has_data(v):
            return [v], None, OK, ""
    # 'Q12-1.' 은 하위 문항인데, 데이터의 Q12_1 은 'Q12.' 의 1) 항목이고 Q12-1 문항은
    # Q12_1_1 에 있을 수 있다. 라벨 맨 앞 문항 번호가 제목과 정확히 같은 변수가
    # 따로 하나 있을 때만 그쪽으로 바꾼다. (한국은행 'Q20-1.' 처럼 Q20_1 의 라벨이
    # 'Q20. 1) …' 이고 다른 후보가 없으면 Q20_1 이 맞다)
    if sub_dash and not k and data.get(code) \
            and _label_code(data.cl[data.get(code)]) not in ("", code.upper()):
        same = [c for c in data.family(code) if c != data.get(code)
                and _label_code(data.cl[c]) == code.upper()
                and data.labels(c) and data.has_data(c)]
        if len(same) == 1:
            return same, None, CHECK, (f"{data.get(code)} 는 다른 문항의 항목이라 라벨 문항 번호가 "
                                       f"맞는 {same[0]} 를 골랐습니다")
    if data.get(code):
        # 'Q4-1. … - 1)' 처럼 항목 번호가 하위 문항 번호를 되풀이한 것이면 그대로 맞다
        if k and not (sub_dash and k == sub_dash[-1]):
            return [data.get(code)], None, CHECK, f"항목 번호 {k} 에 맞는 변수가 없어 {data.get(code)} 를 썼습니다"
        return [data.get(code)], None, OK, ""

    # 라벨로 찾기: 같은 문항 번호 변수 중 제목 뒷부분과 라벨의 항목 문구가 맞는 것
    fam = data.family(code) or data.family(q)
    if not fam:
        return [], None, NEED, f"'{code}' 로 시작하는 변수가 없습니다"
    tail = rest.split("_", 1)[1] if "_" in rest else rest
    tail = re.sub(r"\s-\s.*$", "", tail)
    nt = _norm(tail)
    hits = []
    for c in fam:
        it_txt = _norm(re.sub(r"^\d+\)\s*", "", _item_text(data.cl[c])))
        if it_txt and nt and (it_txt in nt or nt in it_txt):
            hits.append(c)
    if len(hits) > 1:                             # '[가장 높은 환율] 년' — 대괄호 문구로 좁힌다
        nti = _norm(title)
        hits = [c for c in hits
                if all(_norm(x) in nti for x in re.findall(r"\[([^\]]+)\]", data.cl[c]))]
    if len(hits) == 1:
        return hits, None, CHECK, "라벨 문구로 찾았습니다"
    # 대표 변수: 라벨에 '- 세부 지역' 같은 꼬리가 없는 보기형 변수가 하나뿐이면 그것
    main = [c for c in fam if data.labels(c) and not re.search(r"\s-\s*\S", data.cl[c])
            and re.fullmatch(rf"{re.escape(code)}_\d+", c, re.I)]
    if len(main) == 1 and not k:
        return main, None, CHECK, f"같은 문항의 대표 변수 {main[0]} 를 골랐습니다"
    # 복수응답 세트: 문항 번호 바로 아래 변수들(Q1_1 ~ Q1_3)만 본다.
    # 더 깊은 변수(Q1_1_1 = Q1-1 문항)와 섞으면 세트를 못 알아본다.
    ma = data.ma_set(code)
    if ma and not k:
        return ma, "multi", CHECK, "복수응답 세트로 봤습니다"
    labeled = [c for c in fam if data.labels(c)]
    if len(labeled) > 1 and all(data.labels(c) == data.labels(labeled[0]) for c in labeled) \
            and len(labeled) == len(fam) and not k:
        return labeled, "multi", CHECK, "보기가 같은 변수 묶음이라 복수응답으로 봤습니다"
    scored = sorted(((_overlap(tail, data.cl[c]), c) for c in fam), reverse=True)
    if scored and scored[0][0] >= 0.5 and (len(scored) == 1 or scored[0][0] > scored[1][0]):
        return [scored[0][1]], None, CHECK, "라벨이 가장 비슷한 변수를 골랐습니다"
    return [], None, NEED, ("변수가 여러 개라 고르지 못했습니다: "
                            + ", ".join(fam[:6]) + ("…" if len(fam) > 6 else ""))


# ── 베이스 ───────────────────────────────────────────────────────────
def _grid_map(tg: list[tuple], data: Data) -> dict:
    """TG 시트의 'B1-1' 항목 순서 + 항목 문구 → 데이터 변수 (격자형 베이스용)."""
    out, cnt = {}, collections.Counter()
    for r in tg:
        if len(r) < 6:
            continue
        m = re.match(r"\s*([A-Z]+\d+-\d+)\.", _s(r[2]))
        if not (m and _s(r[5])):
            continue
        cnt[m.group(1)] += 1
        base = m.group(1).split("-")[0]
        hit = [c for c in data.numbered(base)
               if data.cl[c].rstrip().endswith(_s(r[5]))]
        if len(hit) == 1:
            out[(m.group(1), cnt[m.group(1)])] = hit[0]
    return out


def _one_cond(raw: str, data: Data, grid: dict, family_hint: list[str]):
    """베이스 조각 하나 → (SPSS 조건, 상태). 못 풀면 (None, NEED)."""
    raw = raw.strip()
    var = r"([A-Za-z]+\d*(?:[-_]\d+)*)"
    m = re.fullmatch(var + r"\s*=\s*(-?\d+)\s*~\s*(-?\d+)", raw)
    if m and data.get(m.group(1)):
        return f"Range({data.get(m.group(1))},{m.group(2)},{m.group(3)})", OK
    m = re.fullmatch(var + r"\s*=\s*(-?\d+(?:\s*(?:,|or)\s*-?\d+)+)", raw, re.I)
    if m and data.get(m.group(1)):
        vals = re.split(r"\s*(?:,|or)\s*", m.group(2), flags=re.I)
        return f"any({data.get(m.group(1))},{','.join(vals)})", OK
    m = re.fullmatch(var + r"\s*(<=|>=|=|<|>)\s*(-?\d+)", raw)
    if m and data.get(m.group(1)):
        return f"{data.get(m.group(1))}{m.group(2)}{m.group(3)}", OK
    # 복수응답 문항: 'Q3=1' 인데 Q3 는 없고 Q3_1 ~ Q3_3 세트 → 보기 1 을 고른 사람
    m = re.fullmatch(var + r"\s*=\s*(\d+(?:\s*(?:,|or)\s*\d+)*)", raw, re.I)
    if m and not data.get(m.group(1)):
        ma = data.ma_set(m.group(1).replace("-", "_"))
        if ma:
            vals = re.split(r"\s*(?:,|or)\s*", m.group(2), flags=re.I)
            parts = [f"any({v},{','.join(ma)})" for v in vals]
            return (parts[0] if len(parts) == 1 else "(" + " | ".join(parts) + ")"), CHECK
    # 격자형: 'B1-1_3)=1' / 'B1-1_1)~3)=1'
    m = re.fullmatch(r"([A-Z]+\d+-\d+)_(\d+)\)(?:\s*~\s*(\d+)\))?\s*=\s*(\d+)", raw)
    if m:
        a, b = int(m.group(2)), int(m.group(3) or m.group(2))
        vs = [grid.get((m.group(1), i)) for i in range(a, b + 1)]
        if all(vs):
            if len(vs) == 1:
                return f"{vs[0]}={m.group(4)}", OK
            return f"any({m.group(4)},{','.join(vs)})", OK
    # 뜻으로 적은 베이스: '변화율=성장' → 같은 문항 변수 중 보기 '성장' 을 가진 것
    m = re.fullmatch(r"(.+?)\s*=\s*(\S.*)", raw)
    if m and family_hint:
        want = _norm(m.group(2))
        for v in family_hint:
            for code, lab in data.labels(v).items():
                if _norm(re.sub(r"^\s*\d+\)\s*", "", str(lab))) == want:
                    return f"{v}={int(code) if float(code).is_integer() else code}", CHECK
    return None, NEED


def _base_cond(raw: str, data: Data, grid: dict, family_hint: list[str]):
    raw = _s(raw)
    if raw.lower() in ("", "all", "전체"):
        return "", OK
    got, worst = [], OK
    for p in re.split(r"\s*&\s*", raw):
        c, s = _one_cond(p, data, grid, family_hint)
        if c is None:
            return None, NEED
        got.append(c)
        worst = CHECK if s == CHECK else worst
    return " & ".join(got), worst


def _versus_vars(title: str, code: str, data: Data):
    """'… - 지방비 vs. 국비 (평균값)' → ([Q16_1_1, Q16_1_2], ['지방비', '국비']).

    'vs' 앞뒤 낱말을 같은 문항 변수의 항목 문구와 맞춘다. 수치형(값 라벨 없음)
    변수가 하나씩 정확히 맞을 때만 돌려준다.
    """
    if not re.search(r"\bvs\.?\b", title, re.I):
        return None
    tail = re.split(r"\s-\s", title)[-1]
    tail = re.sub(r"\([^)]*\)\s*$", "", tail)            # '(평균값)' 제거
    terms = [t.strip() for t in re.split(r"\s*\bvs\.?\s*", tail, flags=re.I) if t.strip()]
    if len(terms) < 2:
        return None
    nums = [c for c in data.family(code) if not data.labels(c)]
    picked = []
    for t in terms:
        hit = [c for c in nums
               if _norm(re.sub(r"^\d+\)\s*", "", _item_text(data.cl[c]))) == _norm(t)]
        if len(hit) != 1:
            return None
        picked.append(hit[0])
    return picked, terms


# ── 한 줄 추론 ─────────────────────────────────────────────────────────
def infer_tables(guide: Guide, df: pd.DataFrame, meta) -> list[TableSpec]:
    data = Data(df, meta)
    grid = _grid_map(guide.tg, data)
    specs: list[TableSpec] = []

    for g in guide.rows:
        title, stats = g.title, g.stats
        m = _CODE.match(title)
        q_code = m.group(1) + m.group(2).replace("-", "_")
        reasons: list[str] = []

        if "summary" in title.lower():
            spec = TableSpec(title, "summary", "", "", row=g.row, stats=stats,
                             base=g.base, extra_banner="")
            specs.append(spec)          # 구성 변수는 아래에서 뒤 줄들을 보고 채운다
            continue

        # 'A vs. B (평균값)' — 같은 문항의 수치형 변수 A·B 를 평균 Summary 로 나란히
        vs_pair = _versus_vars(title, q_code, data)
        if vs_pair:
            cond, cstat = _base_cond(g.base, data, grid, [])
            reasons = [f"'{' vs. '.join(vs_pair[1])}' 를 {', '.join(vs_pair[0])} 의 "
                       "평균 Summary 로 봤습니다"]
            if cond is None:
                reasons.append(f"베이스 '{g.base}' 를 조건으로 바꾸지 못했습니다")
            specs.append(TableSpec(
                title=title, kind="summary", vars=" ".join(v.lower() for v in vs_pair[0]),
                cond=cond or "", sort=False, extra_banner="", stats=stats, base=g.base,
                status=CHECK if cond is not None else NEED,
                reason=" / ".join(reasons), row=g.row))
            continue

        vs, hint, status, why = _find_vars(title, data)
        if why:
            reasons.append(why)
        # '… - 성장' 처럼 베이스가 보기 이름인 줄: 수치형 변수를 쓴다
        fam_hint = []
        if vs:
            fam_hint = data.family(re.sub(r"_\d+$", "", vs[0])) or []
        tail_word = re.search(r"\s-\s*([^\s-][^-]*)$", title)
        if tail_word and re.search(r"[가-힣A-Za-z]", tail_word.group(1)) and "=" in g.base and _norm(tail_word.group(1)) in _norm(g.base):
            fam = data.family(q_code) if not vs else fam_hint
            item = re.search(r"_(\d+)\)", title)
            if item:
                fam = data.family(f"{m.group(1)}_{item.group(1)}")
            nums = [c for c in fam if not data.labels(c)]
            if len(nums) == 1:
                vs, status = nums, CHECK
                reasons = ["베이스가 보기 이름이라 같은 항목의 수치형 변수를 골랐습니다"]
                fam_hint = fam

        if hint:
            kind = hint
        elif not vs:
            kind = "single"
        elif not data.labels(vs[0]):
            kind = "mean"
        elif re.search(r"TOP|BOT", stats, re.I):
            kind = "scale"
        else:
            kind = "single"

        cond, cstat = _base_cond(g.base, data, grid, fam_hint)
        if cond is None:
            miss = [x for x in re.findall(r"[A-Za-z]+\d+(?:[-_]\d+)*", g.base)
                    if not data.get(x) and not re.fullmatch(r"[A-Z]+\d+-\d+", x)]
            reasons.append(f"베이스 '{g.base}' 를 조건으로 바꾸지 못했습니다"
                           + (f" (데이터에 없는 변수: {', '.join(miss)})" if miss else ""))
            cond, status = "", NEED
        elif cstat == CHECK:
            reasons.append("복수응답 문항 베이스를 any(보기, 세트 변수) 로 풀었습니다"
                           if "any(" in cond and not re.search(r"any\([A-Za-z]", cond)
                           else "베이스를 보기 이름으로 풀었습니다")
            status = CHECK if status == OK else status

        extra = ""
        if g.extra_banner:
            extra = data.get(g.extra_banner) or ""
            if not extra:
                reasons.append(f"추가 배너 '{g.extra_banner}' 변수가 없습니다")
                status = NEED
        if g.note and re.search(r"벨류|값으로|리코드|recode|구간", g.note, re.I):
            reasons.append(f"비고 확인: {g.note[:60]}")
            status = CHECK if status == OK else status

        specs.append(TableSpec(
            title=title, kind=kind, vars=data.var_list(vs) if vs else "",
            cond=cond, sort="내림차순" in g.note, extra_banner=extra, stats=stats,
            base=g.base, status=status if vs else NEED,
            reason=" / ".join(reasons), row=g.row))

    # ── Summary 구성: 제목 앞부분이 같은 뒤 줄들의 변수 ──
    for i, sp in enumerate(specs):
        if sp.kind != "summary" or sp.vars:       # 'vs.' 처럼 이미 구성이 정해진 것은 둔다
            continue
        prefix = re.split(r"_(?=[^_]*summary)", sp.title, flags=re.I)[0]
        members = []
        for other in specs[i + 1:]:
            if other.kind == "summary" and other.title.startswith(prefix):
                continue
            if not other.title.startswith(prefix):
                break
            if other.vars and " " not in other.vars:
                members.append(other.vars)
        reason = ""
        if not members:
            m = _CODE.match(sp.title)
            code = m.group(1) + m.group(2).replace("-", "_")
            members = [v.lower() for v in data.numbered(code)]
            reason = f"뒤에 항목 표가 없어 {code}_1 … 전체를 넣었습니다"
        sp.vars = " ".join(members)
        cond, cstat = _base_cond(sp.base, data, grid, [])
        sp.cond = cond or ""
        sp.status = OK if members and not reason else (CHECK if members else NEED)
        sp.reason = reason or ("" if members else "구성 변수를 찾지 못했습니다")
        if cond is None and sp.base.lower() not in ("", "all"):
            sp.reason = (sp.reason + " / " if sp.reason else "") + \
                "베이스는 구성 변수 응답자로 둡니다"
    return specs


# =============================================================================
# 배너
# =============================================================================
@dataclass
class BannerSpec:
    name: str                 # bv1 …
    label: str
    source: str               # 원본 변수
    recode: str = ""          # '(1=1)(2=2)(3 thru hi=3)' — 비우면 COMPUTE
    values: list[tuple[int, str]] = field(default_factory=list)
    status: str = OK
    reason: str = ""
    extra: bool = False       # 추가 배너 (cv 없음)
    members: list[str] = field(default_factory=list)   # 복수응답 배너의 원본 변수들

    def member_names(self, prefix: str = "bv") -> list[str]:
        """복수응답 배너의 bv3_1, bv3_2 … (prefix='cv' 면 cv3_1 …)."""
        base = prefix + self.name[2:]
        return [f"{base}_{i}" for i in range(1, len(self.members) + 1)]


def _mrg_names(banners: list[BannerSpec]) -> dict[str, tuple[str, str]]:
    """복수응답 배너 → (bv 쪽 mrg 이름, cv 쪽 mrg 이름). 첫 번째가 mx11 / mc11."""
    out, j = {}, 0
    for b in banners:
        if b.members:
            j += 1
            out[b.name] = (f"mx1{j}", f"mc1{j}")
    return out


def extra_banners(specs: list[TableSpec], df, meta, start: int) -> list[BannerSpec]:
    """Basic Table 의 '추가배너' 칸에 적힌 변수들로 bv{start}… 를 만든다 (cv 없음)."""
    data = Data(df, meta)
    names = []
    for sp in specs:
        v = data.get(sp.extra_banner) if sp.extra_banner else None
        if v and v not in names:
            names.append(v)
    out = []
    for j, v in enumerate(names, start=start):
        labs = [(int(c), _strip_code(l)) for c, l in sorted(data.labels(v).items())]
        out.append(BannerSpec(f"bv{j}", v, v, "", labs, OK if labs else CHECK,
                              "" if labs else "값 라벨이 없습니다", extra=True))
    return out


def _strip_code(lab) -> str:
    return re.sub(r"^\s*\d+\)\s*", "", str(lab)).strip()


def _numeric_recode(values: list[tuple[int, str]]) -> str | None:
    """'1마리 / 2마리 / 3마리 이상' → '(1=1)(2=2)(3 thru hi=3)'."""
    parts = []
    for code, lab in values:
        nums = [float(x) for x in re.findall(r"\d+(?:\.\d+)?", lab)]
        if re.fullmatch(r"\D*\d+\s*대\D*", lab) and len(nums) == 1:      # '20대'
            parts.append(f"({nums[0]:g} thru {nums[0] + 9:g}={code})")
        elif re.search(r"이상", lab) and len(nums) == 1:
            parts.append(f"({nums[0]:g} thru hi={code})")
        elif re.search(r"이하|미만", lab) and len(nums) == 1:
            hi = nums[0] if "이하" in lab else nums[0] - 1
            parts.append(f"(lo thru {hi:g}={code})")
        elif len(nums) == 2 and re.search(r"[~\-]", lab):
            parts.append(f"({nums[0]:g} thru {nums[1]:g}={code})")
        elif len(nums) == 1:
            parts.append(f"({nums[0]:g}={code})")
        else:
            return None
    return "".join(parts)


def infer_banners(guide: Guide, specs: list[TableSpec], df, meta) -> list[BannerSpec]:
    data = Data(df, meta)
    out: list[BannerSpec] = []
    for i, b in enumerate(guide.banners, start=1):
        name = f"bv{i}"
        want = [_norm(l) for _, l in b.values]
        src = data.get(b.code)
        cands = ([src] if src else []) + [c for c in data.family(b.code.replace("-", "_"))
                                          if c != src]
        chosen, reason, status, recode = None, "", OK, ""
        ma = [] if src else data.ma_set(b.code.replace("-", "_"))
        if ma:
            # 복수응답 배너: bv3_1 = Q1_1 … 로 나눠 만들고 표에서는 /mrg=mx11 로 묶는다
            vals = b.values or [(int(c), _strip_code(l))
                                for c, l in sorted(data.labels(ma[0]).items())]
            spec = BannerSpec(name, b.label, ma[0], "", vals, CHECK, "", members=list(ma))
            spec.reason = (f"복수응답 배너: {', '.join(spec.member_names())} "
                           f"({ma[0]} ~ {ma[-1]}), 표에서는 /mrg 로 묶습니다")
            out.append(spec)
            continue
        for c in cands:                         # 보기 문구가 같은 변수를 찾는다
            have = [_norm(_strip_code(l)) for _, l in sorted(data.labels(c).items())]
            if not want or (have and have == want):
                chosen = c
                break
        if not chosen and want:                 # 보기 개수가 같은 변수 (문구만 조금 다름)
            same_n = [c for c in cands if len(data.labels(c)) == len(want)]
            if same_n:
                chosen = same_n[0]
                reason, status = f"보기 개수가 같은 {chosen} 를 썼습니다 (문구 확인)", CHECK
        if chosen and chosen != src and not reason:
            reason, status = f"보기가 같은 {chosen} 를 썼습니다", CHECK
        if not chosen and src and not data.labels(src):
            rc = _numeric_recode(b.values)
            if rc:
                chosen, recode = src, rc
                reason, status = "숫자 문항이라 보기 문구로 리코드를 만들었습니다", CHECK
        if not chosen and src and data.labels(src):
            # 보기 문구가 같은 코드끼리 잇는다. 안 맞는 것이 있으면 직접 고친다.
            lab = {_norm(_strip_code(l)): int(c) for c, l in data.labels(src).items()}
            pairs = [(lab.get(w), code) for (code, _), w in zip(b.values, want)]
            if all(p[0] is not None for p in pairs):
                chosen = src
                recode = "".join(f"({s}={t})" for s, t in pairs if s != t) or ""
                if recode:
                    recode = "".join(f"({s}={t})" for s, t in pairs)
                    reason, status = "보기 번호가 달라 리코드를 만들었습니다", CHECK
            else:
                chosen, status = src, NEED
                reason = "보기가 원본과 달라 리코드를 직접 적어야 합니다 (예: (1 2=1)(3=2))"
        if not chosen:
            chosen, status = src or "", NEED
            reason = reason or f"'{b.code}' 변수를 찾지 못했습니다"
        values = b.values or [(int(c), _strip_code(l))
                              for c, l in sorted(data.labels(chosen).items())] if chosen else b.values
        out.append(BannerSpec(name, b.label, chosen, recode, values, status, reason))

    return out + extra_banners(specs, df, meta, start=len(out) + 1)


# =============================================================================
# 척도 묶음 · 파생 변수
# =============================================================================
def scale_groups(stats: str, codes: list[float]):
    """'BOT2/SoSo/TOP2' + 코드 → (리코드 문구, 라벨 [(1,'【BOT2】'),…])."""
    n = len(codes)
    bot = re.search(r"BOT\s*(\d+)", stats, re.I)
    top = re.search(r"TOP\s*(\d+)", stats, re.I)
    b = int(bot.group(1)) if bot else 0
    t = int(top.group(1)) if top else 0
    if not (b or t) or b + t > n:
        return None, None
    fmt = lambda c: f"{int(c)}" if float(c).is_integer() else f"{c}"  # noqa: E731
    low, high = codes[:b], codes[n - t:] if t else []
    mid = codes[b:n - t] if t else codes[b:]
    parts, labels = [], []
    if low:
        parts.append(f"({' '.join(map(fmt, low))} = 1)")
        labels.append((1, f"【BOT{b}】"))
    if mid:
        parts.append(f"({' '.join(map(fmt, mid))} = 2)")
        labels.append((2, "【SoSo】"))
    elif low and high:
        labels.append((2, "【SoSo】"))
    if high:
        parts.append(f"({' '.join(map(fmt, high))} = 3)")
        labels.append((3, f"【TOP{t}】"))
    return " ".join(parts), labels


def _fmt(c) -> str:
    return f"{int(c)}" if float(c).is_integer() else f"{c}"


class Derived:
    """Command.sps 에 넣을 파생 변수 줄들을 모은다 (중복 없이, 순서 유지)."""

    def __init__(self):
        self.recode: list[str] = []
        self.vallab: dict[tuple, list[str]] = {}
        self.lines: list[str] = []
        self._seen: set[str] = set()

    def add(self, key: str, line: str, bucket: str = "lines"):
        if key in self._seen:
            return
        self._seen.add(key)
        getattr(self, bucket).append(line)


def _summary_mode(title: str) -> str:
    t = title.lower()
    if "top" in t:
        return "top"
    if "mean" in t or "평균" in t:
        return "mean"
    return "pct"


def prepare(specs: list[TableSpec], df, meta) -> tuple[Derived, list[dict]]:
    """표마다 신텍스에 필요한 조각을 계산한다. (파생 변수, 표별 조각)"""
    data = Data(df, meta)
    der = Derived()
    parts: list[dict] = []
    for sp in specs:
        vs, _bad = parse_vars(sp.vars, data)
        p = {"spec": sp, "vars": vs, "list": data.var_list(vs) if vs else ""}
        if sp.kind == "scale" and vs:
            v = vs[0]
            lv = v.lower()
            codes = data.valid_codes(v)
            rc, labs = scale_groups(sp.stats or "TOP2/SoSo/BOT2", codes)
            if rc:
                der.add(f"r#{lv}", f"Recode {lv} {rc} INTO r#{lv} .", "recode")
                der.vallab.setdefault(tuple(labs), []).append(f"r#{lv}")
            lo, hi = _fmt(codes[0]), _fmt(codes[-1])
            der.add(f"m{lv}", f"If Range({lv},{lo},{hi})  m{lv}={lv} .")
            p.update(points=len(codes), mean=f"m{lv}")
            if re.search(r"100\s*점", sp.stats):
                step = codes[-1] - codes[0]
                der.add(f"mm{lv}", f"If Range({lv},{lo},{hi})  mm{lv}=({lv}-{lo})*(100/{_fmt(step)}) .")
                p["mean100"] = f"mm{lv}"
        elif sp.kind == "mean" and vs:
            lv = vs[0].lower()
            der.add(f"m{lv}", f"Compute m{lv}={lv} .")
            p["mean"] = f"m{lv}"
        elif sp.kind == "summary" and vs:
            mode = _summary_mode(sp.title)
            xs = []
            for v in vs:
                lv = v.lower()
                if mode == "pct":
                    x = f"x{lv}"
                    codes = sorted(data.labels(v)) or [1, 2]
                    yes, rest = _fmt(codes[0]), " ".join(_fmt(c) for c in codes[1:])
                    der.add(x, f"RECODE {lv}  ({yes}=100)({rest}=0) into {x}." if rest
                            else f"RECODE {lv}  ({yes}=100) into {x}.")
                elif mode == "top":
                    codes = data.valid_codes(v)
                    t = re.search(r"TOP\s*(\d+)", sp.title + " " + sp.stats, re.I)
                    t = int(t.group(1)) if t else 2
                    top = codes[len(codes) - t:]
                    x = f"x{''.join(_fmt(c) for c in top)}{lv}"
                    der.add(x, f"recode {lv} ({_fmt(top[0])} thru {_fmt(top[-1])}=100)"
                               f"(sys=sys)(else=0) into {x}.")
                else:
                    x = f"x{lv}"
                    if data.labels(v):
                        codes = data.valid_codes(v)
                        der.add(x, f"if range({lv},{_fmt(codes[0])},{_fmt(codes[-1])}) {x}={lv}.")
                    else:
                        der.add(x, f"Compute {x}={lv}.")
                xs.append((x, _item_text(data.cl.get(v, "")) or v))
            p.update(xvars=xs, mode=mode)
        parts.append(p)
    return der, parts


# =============================================================================
# 신텍스 쓰기
# =============================================================================
def _q(s: str) -> str:
    """SPSS 작은따옴표 문자열 안의 ' 를 겹친다."""
    return str(s).replace("'", "''")


def write_command(banners: list[BannerSpec], specs: list[TableSpec], df, meta, *,
                  work_dir: str, data_file: str, final_file: str) -> str:
    data = Data(df, meta)
    der, _ = prepare(specs, df, meta)
    main = [b for b in banners if not b.extra]
    extra = [b for b in banners if b.extra]
    L = [f'CD "{work_dir}".', "", f"Get file='{data_file}'.", "",
         "*SET PRINTBACK=OFF.", "", "",
         "*_ Banner _____________________________________________________________________.", ""]
    w = max([len(b.name) for b in banners] + [3])
    for b in main:
        if b.status == NEED:
            L.append(f"* 확인 필요: {b.reason}")
        if b.members:                              # 복수응답 배너: bv3_1=Q1_1 …
            for nm, src in zip(b.member_names(), b.members):
                L.append(f"COMPUTE {nm}={src}.")
        elif b.recode:
            L.append(f"RECODE {b.source} {b.recode} into {b.name}.")
        else:
            L.append(f"COMPUTE {b.name.ljust(w)}={b.source}.")
        L.append("")
    if extra:
        L += ["** 추가 배너 .", ""]
        for b in extra:
            L.append(f"Compute {b.name.ljust(w)}={b.source}.")
        L.append("")
    L += ["", "", "*_ Banner by Banner ___________________________________________________________.", ""]
    for b in main:
        if b.members:
            for bv_, cv_ in zip(b.member_names(), b.member_names("cv")):
                L.append(f"Compute {cv_}={bv_}.")
        else:
            L.append(f"Compute {('c' + b.name[1:]).ljust(w)}={b.name.ljust(w)}.")
    # 복수응답 배너의 이름은 표의 /mrg 에 적으므로 변수 라벨을 달지 않는다
    L += ["", "", "", "*_ Variable labels ____________________________________________________________.", "",
          "Variable labels"]
    first = True
    for b in [b for b in main + extra if not b.members]:
        L.append(f"{' ' if first else '/'}{b.name.ljust(w + 1)} '[{_q(b.label)}]'")
        first = False
    for b in [b for b in main if not b.members]:
        L.append(f"/{('c' + b.name[1:]).ljust(w + 1)} '[{_q(b.label)}]'")
    L += [".", "", "", "", "*_ Value labels _______________________________________________________________.", "",
          "Value labels"]
    first = True
    for b in main + extra:
        if b.members:
            pairs = [f"{bv_} {cv_}" for bv_, cv_ in zip(b.member_names(), b.member_names("cv"))]
            L.append(f"{' ' if first else '/'}{pairs[0]}")
            L += [f" {p}" for p in pairs[1:]]
        else:
            L.append(f"{' ' if first else '/'}{b.name if b.extra else f'{b.name} cv{b.name[2:]}'}")
        first = False
        for code, lab in b.values:
            L.append(f"  {str(code).rjust(2) if len(b.values) >= 10 else code}'{_q(lab)}'")
    L += [".", "", "", "", "*_ Banner Check _______________________________________________________________.", ""]
    for b in main:
        if b.members:
            for nm, src in zip(b.member_names(), b.members):
                L.append(f"crosstab {src.ljust(6)} by {nm}.")
        elif b.source:
            L.append(f"crosstab {b.source.ljust(6)} by {b.name}.")
    L += ["", "", "", "*_ RECODE _____________________________________________________________________.", ""]
    for sp in specs:
        if sp.status == NEED:
            L.append(f"* 확인 필요 [{sp.title[:40]}] {sp.reason}")
    L += ["", "", "**_ Top / Bottom ______________________________________________________________.", ""]
    L += der.recode
    if der.vallab:
        L.append("val lab")
        for labs, names in der.vallab.items():
            L.append(f"/{names[0]} to {names[-1]}" if len(names) > 1 else f"/{names[0]}")
            for code, lab in labs:
                L.append(f"{code}'{lab}'")
        L[-1] += "."
    L += der.lines
    L += ["", "", "**_ Mean ______________________________________________________________________.", "",
          "", "", "**_ Summary ___________________________________________________________________.", "",
          "", "", "*______________________________________________________________________________.", ""]
    for c in data.cols:
        L.append(f'var label {c.ljust(22)} "".')
    L += ["", "", "", f"Save outfile='{final_file}'.", "", f"GET FILE='{final_file}'.", "", ""]
    return "\n".join(L)


def _banner_axis(banners: list[BannerSpec], extra: str = ""):
    names = [b.name for b in banners if not b.extra]
    if extra:
        hit = next((b.name for b in banners if b.extra and b.source == extra), None)
        if hit:
            names.append(hit)
    return names


def write_table(specs: list[TableSpec], banners: list[BannerSpec], df, meta, *,
                orientation: str, work_dir: str, final_file: str,
                section_title: str = "응답자 특성") -> str:
    _, parts = prepare(specs, df, meta)
    row = orientation == ROW
    tot, tlab = ("@t3", "■ 전체 ■") if row else ("@t2", "전체")
    head_fmt = "Table format=zero missing('.')" if row else "Table"
    ptotal = (["/ptotal=t2 '사례수'", "/ftotal=t1 '      계'", "/ptotal=t3 '■ 전체 ■'"] if row
              else ["/ptotal=t2 '전체'", "/ftotal=t1 '      계'", "/ptotal=t3 'Base for %'"])
    base_n = "t2" if row else "t3"

    def axes(stub: str, bvs: list[str]) -> str:
        """행 쪽과 열 쪽. stub = 문항 쪽 ('m_down+ t1' 등)."""
        b = "+ ".join(bvs)
        if row:
            return f"/table={tot}+ {b}\n       by t2+ {stub}"
        return f"/table=t3+ {stub}\n       by {tot}+ {b}"

    main = [b for b in banners if not b.extra]
    # 복수응답 배너는 표마다 /mrg=mx11 '[이름]' bv3_1 to bv3_3 로 묶어 배너 자리에 쓴다
    mrg = _mrg_names(banners)
    tok = lambda n: mrg[n][0] if n in mrg else n        # noqa: E731
    mrg_lines = [f"/mrg={mrg[b.name][0]} '[{_q(b.label)}]' "
                 f"{b.member_names()[0]} to {b.member_names()[-1]} " for b in main if b.members]

    def with_mrg(block: list[str]) -> list[str]:
        """'Table …' 줄 바로 아래에 복수응답 배너 /mrg 를 넣는다."""
        if not mrg_lines:
            return block
        i = next(i for i, l in enumerate(block) if l.startswith(head_fmt))
        return block[:i + 1] + mrg_lines + block[i + 1:]

    L = [f'CD "{work_dir}".', "", f"GET FILE='{final_file}'.", "", "SET PRINTBACK=OFF.", "", "",
         "*_ Table ______________________________________________________________________.", "", ""]
    # 배너 표
    bv = [tok(b.name) for b in main]
    cv = [mrg[b.name][1] if b.name in mrg else "c" + b.name[1:] for b in main]
    cv_mrg = [f"/mrg={mrg[b.name][1]} '[{_q(b.label)}]' "
              f"{b.member_names('cv')[0]} to {b.member_names('cv')[-1]} " for b in main if b.members]
    L += [f"Compute {tot}=1.", f"val lab {tot} 1'{tlab}'.", f"{head_fmt}  /* banner */"] \
        + mrg_lines + cv_mrg + ["/"] + ptotal
    if row:
        L += [f"/table={tot}+ {'+ '.join(bv)}", f"       by t2+ {'+ '.join(cv)}",
              f"/statistics=count (t2 (paren5.0) 'Base for %')"]
    else:
        L += [f"/table=t3+ {'+ '.join(cv)}", f"       by {tot}+ {'+ '.join(bv)}",
              f"/statistics=count (t3 (paren5.0) '')"]
    L += [f"            cpct  ( (F4.1) '': {' '.join(bv)}  )", f"/title='{_q(section_title)}'",
          "/", "/", "/caption ''.", "", "", "", ""]

    for p in parts:
        sp: TableSpec = p["spec"]
        vs = p["vars"]
        if not vs:
            L += [f"* 건너뜀 (변수 없음): {sp.title}", "", ""]
            continue
        bvs = [tok(n) for n in _banner_axis(banners, sp.extra_banner)]
        bl = " ".join(bvs)
        lst = p["list"]
        sel_vars = (", ".join(x for x, _ in p["xvars"]) if sp.kind == "summary" and p.get("xvars")
                    else lst)
        cond = f" & ({sp.cond})" if sp.cond else ""
        B = []
        if sp.status == NEED:
            B.append(f"* 확인 필요: {sp.reason}")
        B += ["Temp.", f"Select if nval({sel_vars})>0{cond}.", f"Compute {tot}=1."]
        title = sp.title + (" (중복응답)" if sp.kind == "multi" and "중복" not in sp.title else "")
        tail = ["/", "/", ("/sort= m_down /caption ''." if sp.sort else "/caption ''.")]

        if sp.kind in ("single", "multi"):
            B += [f"val lab {tot} 1'{tlab}'.", f"{head_fmt}  /* {sp.kind} */",
                  f"/mrg=m_down '' {lst} /"] + ptotal
            B.append(axes("m_down+ t1" if sp.kind == "single" else "m_down", bvs))
            B += [f"/statistics=count ({base_n} (paren5.0) '')",
                  f"             cpct  ( m_down (F4.1) '': {bl})",
                  f"/title='{_q(title)}'"] + tail
        elif sp.kind == "mean":
            v, m = lst, p["mean"]
            B += [f"val lab {tot} 1'{'사례수' if row else tlab}'.",
                  f"Variable labels {v} '' .", f"Variable labels {m} '' .",
                  f"{head_fmt}  /* open mean () */", "/", f"/obser={m}"] + ptotal
            B.append(axes(f"{v}+ t1+ {m}", bvs))
            B += [f"/statistics=Count ({base_n} (paren5.0) '')",
                  f"            cpct  ({v} (F4.1) '': {bl} )", "/",
                  f"/statistics=MEAN ({m} (F5.2) '평균')", f"/title='{_q(title)} '"] + tail
        elif sp.kind == "scale":
            v, m, mm = lst, p["mean"], p.get("mean100")
            obs = f"{m} {mm}" if mm else m
            stub = f"{v} + t1+ m_scale+ {m}" + (f"+ {mm}" if mm else "")
            mean_line = (f"/statistics=MEAN  ({m} (F5.2) '[{p['points']}점 평균]')"
                         + (f" MEAN  ({mm} (F5.2) '[100점 환산 평균]')" if mm else ""))
            B += [f"val lab {tot} 1'{tlab}'.", f"{head_fmt}  /* scale ({p['points']}) */",
                  f"/mrg=m_scale '' r#{v}", f"/obser={obs}"] + ptotal
            B.append(axes(stub, bvs))
            B += [f"/statistics=count ({base_n} (paren5.0) '')",
                  f"            cpct  ({v} (F4.1) '': {bl} )",
                  f"            cpct  (m_scale (F4.1) '': {bl} )", "/", mean_line, "/",
                  f"/title='{_q(title)} '"] + tail
        else:                                          # summary
            xs = p.get("xvars") or []
            dec = "F4.1" if p.get("mode") in ("pct", "top") else "F5.2"
            B += [f"val lab {tot} 1'{tlab}'.", f"{head_fmt}  /* top2 summary() */", "//",
                  f"/obser={' '.join(x for x, _ in xs)}"] + ptotal
            B.append(axes("+ ".join(x for x, _ in xs), bvs))
            B.append(f"/statistics=count ({base_n} (paren5.0) '')")
            B += [f"/statistics=MEAN  ({x} ({dec}) '{_q(lab)}')" for x, lab in xs]
            B += [f"/title='{_q(title)} '", "/", "/", "/caption ''."]
        L += with_mrg(B) + ["", "", "", ""]
    return "\n".join(L)


_PLAIN = str.maketrans({"\u00a0": " ", "\u2014": "-", "\u2013": "-", "\u2212": "-",
                        "\u200b": "", "\ufeff": ""})


def encode_sps(text: str, encoding: str) -> tuple[bytes, int]:
    """(바이트, 바뀐 글자 수). CP949 에 없는 글자는 ? 로 바뀐다.

    가이드 제목에 섞여 오는 줄바꿈 없는 공백·대시류는 먼저 흔한 글자로 바꾼다.
    """
    text = text.translate(_PLAIN)
    raw = text.replace("\n", "\r\n").encode(encoding, errors="replace")
    if encoding.lower() in ("cp949", "euc-kr"):
        bad = sum(1 for ch in text if ch != "\n" and not _encodable(ch, encoding))
        return raw, bad
    return raw, 0


def _encodable(ch: str, enc: str) -> bool:
    try:
        ch.encode(enc)
        return True
    except UnicodeEncodeError:
        return False
